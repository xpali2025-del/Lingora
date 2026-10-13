import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CONTENT = ROOT / 'content'
OUTPUT = ROOT / 'docs' / 'Lingora_SEO_Content_Draft.xlsx'

records = []
for locale in ('ar', 'en'):
    for path in sorted((CONTENT / locale).glob('*.json')):
        records.append(json.loads(path.read_text(encoding='utf-8')))
product_facts = json.loads((CONTENT / 'product-facts.json').read_text(encoding='utf-8'))
public_drafts = [r for r in records if r['publicationState'] != 'research-only']
indexable_records = [r for r in public_drafts if r['indexable']]
noindex_records = [r for r in public_drafts if not r['indexable']]
research_only = [r for r in records if r['publicationState'] == 'research-only']


def visible_text(record):
    parts = [record['h1'], record['intro']]
    for section in record['sections']:
        parts.append(section['heading'])
        parts.extend(section['paragraphs'])
    for item in record['faqs']:
        parts.extend((item['question'], item['answer']))
    for option in record.get('paymentOptions', []):
        parts.extend((option['group'], option['label'], option['value'], option['details']))
    if record['pageType'] == 'plans':
        facts = product_facts[record['locale']]
        parts.extend((facts['heading'], *facts['facts'], facts['note']))
    return ' '.join(parts)


def word_count(record):
    return len(visible_text(record).split())


wb = Workbook()
readme = wb.active
readme.title = 'Start Here'
for row in [
    ['Lingora SEO & Page Copy — DRAFT'],
    ['Scope', f'{len(records)} content records; {len(public_drafts)} served draft routes; {len(noindex_records)} noindex,nofollow routes; {len(indexable_records)} owner-approved indexable About routes; {len(research_only)} country records withheld.'],
    ['Markets', 'The 12 country pages are research-only, excluded from static routes, and must not be restored by swapping country names into generic copy.'],
    ['Search Console', 'No export was supplied. Clicks, impressions, CTR, and position remain blank; CSV imports do not auto-assign keyword owners.'],
    ['Keyword tools', 'Search volume and difficulty remain blank until actual dated exports are provided.'],
    ['Legal/commercial', 'Policy copy and operational terms remain blocked pending actual data flows, owner-approved prices/fulfilment/refunds, and legal review.'],
    ['Product claims', 'Super feature summary is maintained once in content/product-facts.json, links to Duolingo, and must be rechecked before publication.'],
    ['Publication', 'Only the bilingual About pages are indexable at the owner’s request; all other served routes remain noindex,nofollow drafts. Country-market routes remain withheld. Google indexing is not guaranteed.'],
    ['Completeness', 'Word targets are project editorial ranges, not a Google ranking formula. Resolve user tasks with evidence; do not pad to a number.'],
    ['Duplicate/cannibalization gate', 'Checks active keyword overlap, metadata/FAQ uniqueness, and five-word shingle similarity; country route withholding prevents doorway pages.'],
    ['Commands', 'npm run seo:check | npm run seo:gsc | npm run check:all'],
    ['Research plan', 'See docs/SEO_INTENT_RESEARCH_PLAN.md and docs/KEYWORD_RESEARCH_WORKFLOW.md.'],
]:
    readme.append(row)

page = wb.create_sheet('Page Map')
page.append([
    'Locale', 'URL', 'Publication state', 'Indexing state', 'Page type', 'Intent stage',
    'Primary keyword (proposed)', 'Current words', 'Target words', 'Title',
    'Description', 'H1', 'Volume / month', 'Difficulty', 'GSC clicks',
    'GSC impressions', 'GSC CTR', 'GSC position', 'Research status',
    'Publish blockers',
])
intent = wb.create_sheet('Intent Map')
intent.append([
    'Locale', 'URL', 'State', 'Search stage', 'Primary user need',
    'Page job / promised answer', 'Distinct from / cannibalization guard',
    'Intent research status', 'Primary query (proposed)',
])
keywords = wb.create_sheet('Keyword Clusters')
keywords.append([
    'Locale', 'URL', 'Page state', 'Cluster type', 'Keyword phrase',
    'Active primary owner?', 'Search volume', 'Difficulty', 'GSC clicks',
    'GSC impressions', 'GSC CTR', 'GSC position', 'Research status',
    'Source / next step',
])
copy = wb.create_sheet('Page Copy')
copy.append([
    'Locale', 'URL', 'Publication state', 'H1', 'Introduction',
    'Section heading', 'Draft paragraphs', 'Target words', 'Current words',
    'Editorial state',
])
faq = wb.create_sheet('FAQs')
faq.append(['Locale', 'URL', 'Publication state', 'Question', 'Draft answer', 'Status'])
payments = wb.create_sheet('Payment Methods')
payments.append(['Locale', 'URL', 'Market group', 'Method', 'Owner-supplied value', 'Instructions / verification status'])
checks = wb.create_sheet('Checks & Blockers')
checks.append(['Locale', 'URL', 'State', 'Check', 'Result / note'])
owner_inputs = wb.create_sheet('Owner Inputs')
owner_inputs.append(['Blocking field', 'What is needed to finish the page truthfully', 'Owner answer', 'Supporting source / URL', 'Reviewed by', 'Review date', 'Status'])
for row in [
    ['Business identity', 'Legal entity name, registered jurisdiction, public contact route, and operator approval.', 'Public phone and emails supplied: +966554751363; info@lingora.com; sales@lingora.com. Legal entity and delivery of these channels are not independently verified.', 'Owner confirmation in chat, 2026-10-13', 'Owner', '2026-10-13', 'PARTIAL — LEGAL ID / CONTACT DELIVERY UNVERIFIED'],
    ['About page contact', 'Owner-approved public contact details for the indexable About pages.', 'Phone: +966554751363; general: info@lingora.com; sales: sales@lingora.com.', 'Owner confirmation in chat, 2026-10-13', 'Owner', '2026-10-13', 'OWNER-SUPPLIED — DELIVERY NOT TESTED'],
    ['Resale rights / source', 'Written proof of the right to supply or resell the redeem links, its scope, and current supplier/source.', None, None, None, None, 'BLOCKED'],
    ['Redeem terms', 'Exact product/tier, whether individual or Family, duration start/end, expiry, eligibility, and tested redemption steps.', 'Owner confirms individual and Family Duolingo Super; Family is not discounted; Duolingo Max is not offered. Duration, eligibility, supply rights, and redemption steps remain unverified.', 'Owner confirmation in chat, 2026-10-13', 'Owner', '2026-10-13', 'PARTIAL — TERMS / RIGHTS UNVERIFIED'],
    ['Price / tax', 'Owner-approved price and currency by track, treatment of taxes/fees, and checkout amount shown to the buyer.', None, None, None, None, 'BLOCKED'],
    ['Payment availability', 'Owner supplied Yemen and Gulf methods for the shared payment page; independently test merchant receiving, currency, fees, settlement, and market access before checkout.', 'See the Payment Methods sheet; details are not independently validated.', 'User-provided payment note, 2026-10-13', None, '2026-10-13', 'OWNER-SUPPLIED — VERIFY'],
    ['BEP20 wallet address', 'Supply and test the complete BNB Smart Chain wallet address. The string supplied in chat is incomplete and is not used.', None, None, None, None, 'BLOCKED — ADDRESS INCOMPLETE'],
    ['Payment API credential', 'Revoke the exposed credential and issue a replacement. Store only in server-side deployment secrets after API provider documentation and webhook signing are confirmed.', None, None, None, None, 'BLOCKED — ROTATE SECRET'],
    ['Payment gateway integration', 'Provide provider documentation, endpoint/webhook details, signature verification procedure, and a successful low-value test in a backend; never embed a secret in static HTML/JS.', None, None, None, None, 'BLOCKED — NO VERIFIED API DOCS'],
    ['Exchange rates / transaction fees', 'Owner-approved display currency, conversion source/time, network fees, settlement and refund handling.', None, None, None, None, 'BLOCKED'],

    ['Fulfilment / stock', 'Real delivery route, stock controls, out-of-stock handling, and tested paid-but-unfulfilled remedy.', None, None, None, None, 'BLOCKED'],
    ['Support', 'Working contact channel, responsible operator, tested escalation route, and approved response-time statement.', None, None, None, None, 'BLOCKED'],
    ['Refund / cancellation', 'Owner-approved rules for invalid links, failed delivery, cancellation, chargeback, and refund timing; legal review.', None, None, None, None, 'BLOCKED'],
    ['Privacy / data flow', 'Actual fields collected, hosting/payment/support processors, log/analytics data, retention, deletion, and approved privacy text.', None, None, None, None, 'BLOCKED'],
    ['Market page evidence', 'For each country: a distinct user need plus dated local payment/consumer/tax/support evidence and original answer—or keep merged.', None, None, None, None, 'BLOCKED'],
    ['Keyword research', 'Dated keyword-tool exports by language and actual target market, with provider and matching methodology.', None, None, None, None, 'WAITING FOR EXPORT'],
    ['Search Console', 'queries.csv, pages.csv, and preferably query-page.csv from the verified property.', None, None, None, None, 'WAITING FOR EXPORT'],
]:
    owner_inputs.append(row)

tools = wb.create_sheet('Tools & GSC')
for row in [
    ['Tool / input', 'Use', 'Metric or evidence', 'Availability / note'],
    ['Google Keyword Planner', 'Validate candidate terms by language and actual market.', 'Average monthly searches; advertiser competition is not organic SEO difficulty.', 'No export supplied. Record provider, date, geo and match type.'],
    ['Google Trends', 'Compare relative interest and seasonality.', 'Relative index, not absolute volume.', 'Supporting evidence only.'],
    ['Google Search Console', 'Measure actual organic queries and pages after property verification.', 'Clicks, impressions, CTR, average position.', 'No connection/export. Upload CSVs.'],
    ['Optional Ahrefs / Semrush', 'Supplement keyword research.', 'Provider-specific demand/difficulty estimates.', 'Keep provider/date; never present estimates as actual GSC metrics.'],
    ['Search Console CSV importer', 'Review queries.csv, pages.csv, optional query-page.csv.', 'Creates reports/search-console/summary.md for the 14 served draft routes.', 'Run npm run seo:gsc; no copy or owner is assigned automatically.'],
    ['Near-duplicate gate', 'Compare same-language copy via five-word shingles.', 'FAIL ≥25% containment; WATCH ≥15%.', 'Run npm run seo:check.'],
    ['Cannibalization gate', 'Check exact ownership and high token overlap between active multiword primaries.', 'Overlap warnings apply only to active draft routes.', 'Run npm run seo:check.'],
    ['Market publication gate', 'Require genuinely unique local evidence before restoring any market URL.', 'Source + date + useful local answer + original copy.', 'All 12 country records are currently research-only.'],
]:
    tools.append(row)

for record in records:
    count = word_count(record)
    state = record['publicationState']
    page.append([
        record['locale'], record['route'], state,
        'indexable' if record['indexable'] else 'noindex,nofollow',
        record['pageType'], record['searchIntent']['stage'],
        record['keywordResearch']['primaryKeyword'],
        count, record['targetWords'], record['title'], record['description'],
        record['h1'], None, None, None, None, None, None,
        record['searchIntent']['researchStatus'], '; '.join(record['publishBlockers']),
    ])
    intent.append([
        record['locale'], record['route'], state, record['searchIntent']['stage'],
        record['searchIntent']['userNeed'], record['searchIntent']['pageJob'],
        record['searchIntent']['distinctFrom'], record['searchIntent']['researchStatus'],
        record['keywordResearch']['primaryKeyword'],
    ])
    for kind, values in [
        ('Primary', [record['keywordResearch']['primaryKeyword']]),
        ('Short-tail', record['keywordResearch']['shortTail']),
        ('Long-tail', record['keywordResearch']['longTail']),
    ]:
        for phrase in values:
            active_owner = (
                'Yes' if kind == 'Primary' and state != 'research-only' else 'No'
            )
            keywords.append([
                record['locale'], record['route'], state, kind, phrase,
                active_owner, None, None, None, None, None, None,
                record['keywordResearch']['status'], record['keywordResearch']['source'],
            ])
    for section in record['sections']:
        copy.append([
            record['locale'], record['route'], state, record['h1'], record['intro'],
            section['heading'], '\n\n'.join(section['paragraphs']),
            record['targetWords'], count, record['status'],
        ])
    if record['pageType'] == 'plans':
        facts = product_facts[record['locale']]
        copy.append([
            record['locale'], record['route'], state, record['h1'], record['intro'],
            facts['heading'], '\n\n'.join([*facts['facts'], facts['note']]),
            record['targetWords'], count, 'Central source-backed product facts',
        ])
    for item in record['faqs']:
        faq.append([
            record['locale'], record['route'], state,
            item['question'], item['answer'], record['status'],
        ])
    for option in record.get('paymentOptions', []):
        payments.append([
            record['locale'], record['route'], option['group'], option['label'],
            option['value'], option['details'],
        ])
    if state == 'research-only':
        checks.append([
            record['locale'], record['route'], state, 'Route safety',
            'Withheld from the public export until distinct local evidence is approved.',
        ])
    else:
        checks.append([
            record['locale'], record['route'], state, 'Content completeness',
            f'{count}/{record["targetWords"]} words; this is an editorial warning, not a ranking target. Resolve the task with verified evidence; never add filler.',
        ])
    checks.append([
        record['locale'], record['route'], state, 'Keyword metrics',
        'Volume, difficulty and GSC fields intentionally blank pending real source exports.',
    ])
    for blocker in record['publishBlockers']:
        checks.append([
            record['locale'], record['route'], state, 'Publication blocker', blocker,
        ])

header_fill = PatternFill('solid', fgColor='355A43')
header_font = Font(color='FFFFFF', bold=True)
alt_fill = PatternFill('solid', fgColor='F2F6F1')
thin = Side(style='hair', color='DDE5DD')
for sheet in wb.worksheets:
    sheet.freeze_panes = 'A2'
    if sheet.title != 'Start Here':
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        sheet.row_dimensions[1].height = 38
    for row_index in range(2, sheet.max_row + 1):
        if row_index % 2 == 0:
            for cell in sheet[row_index]:
                cell.fill = alt_fill
        for cell in sheet[row_index]:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = Border(bottom=thin)
    for col_index in range(1, sheet.max_column + 1):
        values = [
            str(sheet.cell(row_index, col_index).value or '')
            for row_index in range(1, sheet.max_row + 1)
        ]
        width = min(
            max(12, max((len(value.split('\n')[0]) for value in values), default=12) + 2),
            54,
        )
        if sheet.title in ('Intent Map', 'Page Copy', 'FAQs', 'Checks & Blockers', 'Start Here') and col_index > 2:
            width = max(width, 40)
        sheet.column_dimensions[get_column_letter(col_index)].width = width
    sheet.sheet_view.showGridLines = False

readme.column_dimensions['A'].width = 30
readme.column_dimensions['B'].width = 112
for row_index in range(1, readme.max_row + 1):
    readme.cell(row_index, 1).font = Font(bold=True, color='355A43')
    readme.cell(row_index, 2).alignment = Alignment(wrap_text=True, vertical='top')
readme.row_dimensions[1].height = 30

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)
print(
    f'Wrote {OUTPUT.relative_to(ROOT)}: {len(records)} records, '
    f'{len(public_drafts)} active draft routes, {len(research_only)} research-only markets.'
)
